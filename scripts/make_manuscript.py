"""Build the manuscript (.docx) from docs/manuscript/manuscript.md.

Supported Markdown subset (enough for a research paper):
  # Title                        -> paper title
  Authors: ... / Affiliation: ...-> centered front-matter lines
  ## Heading / ### Sub-heading   -> section / subsection headings
  paragraphs with **bold**, *italic*, `code`
  - bullets, 1. numbered items
  Table: caption  + a pipe table -> captioned table
  ![Fig. N. caption](path.png)   -> centered figure with caption
  $$ formula $$                  -> centered display equation (plain text)
  [n] reference text             -> reference list entry (hanging indent)

usage: python scripts/make_manuscript.py   -> docs/manuscript/BO_Compiler_Pass_Selection_Manuscript.docx
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "docs" / "manuscript" / "manuscript.md"
OUT = ROOT / "docs" / "manuscript" / "BO_Compiler_Pass_Selection_Manuscript.docx"
FONT = "Times New Roman"


AUTHORS = ROOT / "docs" / "manuscript" / "authors.yaml"


def apply_authors(lines: list[str]) -> list[str]:
    """Fill the Authors / Affiliation / Contact lines from docs/manuscript/authors.yaml."""
    if not AUTHORS.exists():
        return lines
    import yaml
    a = yaml.safe_load(AUTHORS.read_text(encoding="utf-8")) or {}
    members = [str(m) for m in a.get("members", []) if str(m).strip()]
    affiliation = ", ".join(str(a[k]) for k in ("course", "department", "university") if a.get(k))
    out = []
    for line in lines:
        if line.startswith("Contact:"):
            continue                      # regenerated below from authors.yaml
        if line.startswith("Authors:") and members:
            line = "Authors: " + ", ".join(members)
        elif line.startswith("Affiliation:") and affiliation:
            out.append("Affiliation: " + affiliation)
            if a.get("email"):
                out.append("Contact: " + str(a["email"]))
            continue
        out.append(line)
    return out


def set_cell_shading(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def add_runs(par, text: str, size: float | None = None, italic_all: bool = False) -> None:
    """Inline **bold**, *italic* and `code`."""
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = par.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("*"):
            r = par.add_run(tok[1:-1]); r.italic = True
        elif tok.startswith("`"):
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 10.5) - 1)
            continue
        else:
            r = par.add_run(tok)
        if italic_all:
            r.italic = True
        if size:
            r.font.size = Pt(size)


def add_page_number(section) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)


def build() -> int:
    if not SRC.exists():
        print(f"{SRC} not found")
        return 1
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(1))
    sec.top_margin = sec.bottom_margin = Inches(0.9)
    add_page_number(sec)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(10.5)
    st.paragraph_format.space_after = Pt(5)
    st.paragraph_format.line_spacing = 1.12
    for name, size in (("Heading 1", 12), ("Heading 2", 11)):
        h = doc.styles[name]
        h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = RGBColor(0, 0, 0)
        rf = h.element.rPr.rFonts
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(att), None)          # theme fonts would override FONT
        for att in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(att), FONT)
        h.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        h.paragraph_format.space_after = Pt(4)
    if "Heading 2" in [s.name for s in doc.styles]:
        doc.styles["Heading 2"].font.italic = True

    lines = apply_authors(SRC.read_text(encoding="utf-8").splitlines())
    i, table_caption, in_abstract = 0, None, False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:].strip()); r.bold = True; r.font.size = Pt(18)
            p.paragraph_format.space_after = Pt(10)
        elif re.match(r"^(Authors|Affiliation|Contact):", line):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, line.split(":", 1)[1].strip(), size=11 if line.startswith("Authors") else 10,
                     italic_all=not line.startswith("Authors"))
            p.paragraph_format.space_after = Pt(2)
        elif line.startswith("## "):
            title = line[3:].strip()
            in_abstract = title.lower() in ("abstract", "index terms")
            if in_abstract:
                p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
                r = p.add_run(title); r.bold = True; r.italic = True
            else:
                doc.add_heading(title, level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif line.startswith("Table:"):
            table_caption = line[6:].strip()
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if table_caption:
                cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cr = cp.add_run(table_caption); cr.font.size = Pt(9.5); cr.bold = True
                cp.paragraph_format.keep_with_next = True
                table_caption = None
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ri, row in enumerate(rows):
                for ci in range(ncol):
                    cell = t.cell(ri, ci)
                    cell.text = ""
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    add_runs(par, row[ci] if ci < len(row) else "", size=8.5)
                    if ri == 0:
                        for r in par.runs:
                            r.bold = True
                        set_cell_shading(cell, "E8EDF5")
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        elif (m := re.match(r"^!\[(.*?)\]\((.*?)\)", line)):
            img = (ROOT / m.group(2)).resolve()
            if img.exists():
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(img), width=Inches(6.2))
                p.paragraph_format.keep_with_next = True
            cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(cp, m.group(1), size=9.5)
            cp.paragraph_format.space_after = Pt(10)
        elif line.startswith("$$"):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line.strip("$ ").strip()); r.italic = True; r.font.size = Pt(11.5)
        elif (m := re.match(r"^\s*[-*]\s+(.*)", line)):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, m.group(1))
            p.paragraph_format.space_after = Pt(2)
        elif (m := re.match(r"^\s*\d+\.\s+(.*)", line)):
            p = doc.add_paragraph(style="List Number"); add_runs(p, m.group(1))
            p.paragraph_format.space_after = Pt(2)
        elif re.match(r"^\[\d+\]", line):
            p = doc.add_paragraph(); add_runs(p, line, size=9.5)
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_runs(p, line, italic_all=False)
            if in_abstract:
                for r in p.runs:
                    r.font.size = Pt(10)
        i += 1
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(build())
